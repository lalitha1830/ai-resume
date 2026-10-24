"""
services/resume_parser.py — Resume text extraction.

This module is responsible for reading PDF, DOCX, and JSON files and returning
their plain-text content.  It is completely separate from Flask routes
so it can be tested independently.

Supported formats:
  • PDF  — using PyMuPDF (the 'fitz' library)
  • DOCX — using python-docx
  • JSON — structured data (LinkedIn exports, custom formats, etc.)

Important notes for beginners:
  • Some PDFs are scanned images.  They look like documents but contain
    no machine-readable text.  PyMuPDF cannot read those — OCR would be
    needed.  We detect this case and tell the user.
  • No file size limit is enforced here.
  • We never print or log the resume content for privacy reasons.
  • Uploaded files are handled as file-like objects (not saved to disk).
"""

import io
import json as json_module
import logging
from typing import Optional

# PyMuPDF is imported as 'fitz' (historical name)
try:
    import fitz  # PyMuPDF
    PYMUPDF_AVAILABLE = True
except ImportError:
    PYMUPDF_AVAILABLE = False

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

logger = logging.getLogger(__name__)

# Minimum number of characters we consider "readable" text
MIN_TEXT_LENGTH = 50


class ResumeParseError(Exception):
    """Raised when we cannot extract usable text from a resume."""
    pass


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extract all text from a PDF file given as raw bytes.

    Args:
        file_bytes: The raw binary content of the PDF file.

    Returns:
        Extracted text as a single string.

    Raises:
        ResumeParseError: If the PDF cannot be opened, is empty, or is
                          a scanned image with no machine-readable text.
    """
    if not PYMUPDF_AVAILABLE:
        raise ResumeParseError(
            "PDF processing library (PyMuPDF) is not installed. "
            "Run: pip install PyMuPDF"
        )

    try:
        # Open the PDF from memory — we never write to disk
        pdf_document = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as exc:
        logger.warning("Failed to open PDF: %s", type(exc).__name__)
        raise ResumeParseError(
            "We couldn't open this PDF file. It may be corrupted or password-protected."
        ) from exc

    if pdf_document.page_count == 0:
        pdf_document.close()
        raise ResumeParseError("This PDF appears to have no pages.")

    # Extract text from every page and join with newlines
    all_text_parts = []
    for page_number in range(pdf_document.page_count):
        try:
            page = pdf_document.load_page(page_number)
            page_text = page.get_text("text")  # plain-text extraction
            if page_text:
                all_text_parts.append(page_text.strip())
        except Exception as exc:
            logger.warning(
                "Could not read page %d: %s", page_number + 1, type(exc).__name__
            )
            # Continue to the next page instead of failing entirely

    pdf_document.close()

    full_text = "\n\n".join(all_text_parts).strip()

    if len(full_text) < MIN_TEXT_LENGTH:
        raise ResumeParseError(
            "We couldn't extract readable text from this PDF. "
            "It may be a scanned image. "
            "Please use a PDF that contains machine-readable text, or convert your resume to DOCX."
        )

    return full_text


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract all text from a DOCX file given as raw bytes.

    Args:
        file_bytes: The raw binary content of the DOCX file.

    Returns:
        Extracted text as a single string.

    Raises:
        ResumeParseError: If the file is corrupted or contains no readable text.
    """
    if not DOCX_AVAILABLE:
        raise ResumeParseError(
            "DOCX processing library (python-docx) is not installed. "
            "Run: pip install python-docx"
        )

    try:
        # Open the DOCX from memory — we never write to disk
        doc = DocxDocument(io.BytesIO(file_bytes))
    except Exception as exc:
        logger.warning("Failed to open DOCX: %s", type(exc).__name__)
        raise ResumeParseError(
            "We couldn't open this DOCX file. It may be corrupted or not a valid Word document."
        ) from exc

    # Each paragraph in a Word document is a separate text block
    paragraphs = [para.text.strip() for para in doc.paragraphs if para.text.strip()]

    # Also extract text from tables (resume skills, education tables, etc.)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text and cell_text not in paragraphs:
                    paragraphs.append(cell_text)

    full_text = "\n".join(paragraphs).strip()

    if len(full_text) < MIN_TEXT_LENGTH:
        raise ResumeParseError(
            "We couldn't extract readable text from this Word document. "
            "Please ensure the document contains text (not just images)."
        )

    return full_text


def extract_text_from_json(file_bytes: bytes) -> str:
    """
    Extract readable text from a JSON resume file.

    Accepts any JSON format — LinkedIn data exports, custom schemas, etc.
    All string values are recursively collected and joined into one block
    of text that the AI can analyse.

    Args:
        file_bytes: The raw binary content of the JSON file.

    Returns:
        Extracted text as a single string.

    Raises:
        ResumeParseError: If the file is not valid JSON or contains no text.
    """
    try:
        text = file_bytes.decode("utf-8", errors="replace")
        data = json_module.loads(text)
    except (json_module.JSONDecodeError, ValueError) as exc:
        logger.warning("Failed to parse JSON resume: %s", exc)
        raise ResumeParseError(
            "The uploaded file does not appear to be valid JSON. "
            "Please upload a properly formatted .json resume."
        ) from exc

    # Recursively collect all string values from any nested structure
    collected: list[str] = []

    def _collect(node: object) -> None:
        if isinstance(node, str):
            stripped = node.strip()
            if stripped:
                collected.append(stripped)
        elif isinstance(node, dict):
            for v in node.values():
                _collect(v)
        elif isinstance(node, list):
            for item in node:
                _collect(item)

    _collect(data)

    full_text = "\n".join(collected).strip()

    if len(full_text) < MIN_TEXT_LENGTH:
        raise ResumeParseError(
            "We couldn't extract readable text from this JSON file. "
            "Please make sure the JSON contains resume content as string values."
        )

    return full_text


def parse_resume(file_bytes: bytes, filename: str) -> str:
    """
    Main entry point: detect the file type and extract text.

    Supported extensions: .pdf, .docx, .json
    No file size limit is enforced here.

    Args:
        file_bytes: Raw bytes of the uploaded file.
        filename:   Original filename (used to detect the extension).

    Returns:
        Extracted plain text from the resume.

    Raises:
        ResumeParseError: If the file type is unsupported or extraction fails.
    """
    if not file_bytes:
        raise ResumeParseError("The uploaded file is empty.")

    lower_name = filename.lower().strip()

    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif lower_name.endswith(".json"):
        return extract_text_from_json(file_bytes)
    else:
        raise ResumeParseError(
            "Unsupported file type. Please upload a PDF (.pdf), "
            "Word document (.docx), or JSON file (.json)."
        )
